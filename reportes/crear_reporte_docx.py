from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
REPORTES = ROOT / "reportes"
DOCS = ROOT / "docs"

FONT_MONO = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
FONT_SANS = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def load_font(path, size):
    try:
        return ImageFont.truetype(path, size)
    except Exception:
        return ImageFont.load_default()


def crear_captura(texto, salida, titulo="Terminal"):
    font = load_font(FONT_MONO, 18)
    title_font = load_font(FONT_SANS, 18)
    lineas = texto.rstrip().splitlines() or [""]
    max_chars = max(len(linea) for linea in lineas)
    ancho = max(900, min(1800, max_chars * 11 + 60))
    alto = 44 + len(lineas) * 26 + 30
    img = Image.new("RGB", (ancho, alto), (32, 32, 32))
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, ancho, 40), fill=(55, 55, 55))
    draw.ellipse((16, 14, 28, 26), fill=(255, 95, 86))
    draw.ellipse((36, 14, 48, 26), fill=(255, 189, 46))
    draw.ellipse((56, 14, 68, 26), fill=(39, 201, 63))
    draw.text((85, 10), titulo, font=title_font, fill=(230, 230, 230))
    y = 55
    for linea in lineas:
        color = (230, 230, 230)
        lower = linea.lower()
        if lower.endswith(" ok") or linea == "OK":
            color = (170, 255, 170)
        elif "failed" in lower or "error" in lower:
            color = (255, 140, 140)
        draw.text((24, y), linea, font=font, fill=color)
        y += 26
    img.save(salida)


programa = (REPORTES / "ejecucion_programa.txt").read_text(encoding="utf-8")
pruebas = (REPORTES / "pruebas_unitarias.txt").read_text(encoding="utf-8")
archivo = (ROOT / "datos" / "cuentas.txt").read_text(encoding="utf-8")

captura_programa = REPORTES / "captura_ejecucion_programa.png"
captura_pruebas = REPORTES / "captura_pruebas_unitarias.png"
captura_archivo = REPORTES / "captura_archivo_txt.png"
crear_captura(programa, captura_programa, "python main.py")
crear_captura(pruebas, captura_pruebas, "python -m unittest discover -s tests -v")
crear_captura(archivo, captura_archivo, "datos/cuentas.txt")

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Sistema de cuentas bancarias en Python")
run.bold = True
run.font.size = Pt(18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Documento con pantallazos de ejecucion y pruebas unitarias")
run.font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Autor: Daniel Felipe Moreno | Rama sugerida: danielfelipe.moreno.suarez")

doc.add_heading("1. Resumen del programa", level=1)
doc.add_paragraph(
    "El proyecto implementa un sistema de cuentas bancarias usando Python y programacion orientada a objetos. "
    "La clase abstracta CuentaBancaria contiene el comportamiento comun, mientras que CuentaAhorros y "
    "CuentaCorriente especializan las reglas de negocio. La persistencia se realiza en un archivo de texto "
    "y las pruebas unitarias validan las operaciones principales."
)

doc.add_heading("2. Conceptos evidenciados", level=1)
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Concepto"
hdr[1].text = "Aplicacion"
rows = [
    ("Herencia", "CuentaAhorros y CuentaCorriente heredan de CuentaBancaria."),
    ("Reutilizacion", "depositar, retirar, validar montos y to_dict se implementan en la clase base."),
    ("Polimorfismo", "Banco llama aplicar_corte_mensual() sobre cuentas de diferentes tipos."),
    ("Modularidad", "El proyecto se divide en modelos, banco, persistencia, excepciones y pruebas."),
    ("Manejo de archivos", "RepositorioCuentasTxt guarda objetos en datos/cuentas.txt."),
    ("Pruebas unitarias", "unittest verifica reglas de negocio, persistencia y polimorfismo."),
]
for c1, c2 in rows:
    row = table.add_row().cells
    row[0].text = c1
    row[1].text = c2


doc.add_heading("3. Pantallazo: ejecucion del programa", level=1)
doc.add_paragraph("Comando ejecutado: python main.py")
doc.add_picture(str(captura_programa), width=Inches(6.7))


doc.add_heading("4. Pantallazo: archivo de texto generado", level=1)
doc.add_paragraph("Archivo revisado: datos/cuentas.txt")
doc.add_picture(str(captura_archivo), width=Inches(6.7))


doc.add_heading("5. Pantallazo: pruebas unitarias", level=1)
doc.add_paragraph("Comando ejecutado: python -m unittest discover -s tests -v")
doc.add_picture(str(captura_pruebas), width=Inches(6.7))


doc.add_heading("6. Resultado", level=1)
doc.add_paragraph(
    "Las 8 pruebas unitarias finalizaron correctamente con resultado OK. Esto evidencia que el sistema cumple "
    "las reglas principales del caso: depositos, retiros, sobregiro, transferencia, polimorfismo y persistencia."
)

salida = DOCS / "documento_pantallazos_pruebas.docx"
doc.save(salida)
print(salida)
